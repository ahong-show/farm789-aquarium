import streamlit as st
import pandas as pd
import sqlite3
import datetime
import io
import os
import plotly.express as px

# Library Excel Styling
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# ==========================================
# 1. KONFIGURASI & CUSTOM CSS TEMA TERANG
# ==========================================
st.set_page_config(
    page_title="FARM 789 AQUARIUM - Keuangan",
    page_icon="🐠",
    layout="wide",
    initial_sidebar_state="expanded"
)

# CSS Custom Khusus Tema Terang Bersih dengan Border Input Sangat Jelas
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;600;700;800&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    /* Background Utama Aplikasi Putih Bersih */
    .stApp {
        background-color: #ffffff !important;
    }

    /* --- BORDER INPUT DIPERJELAS --- */
    /* Input Text, Number, Date, Selectbox */
    div[data-baseweb="input"], 
    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        border: 2px solid #94a3b8 !important;
        border-radius: 8px !important;
        transition: all 0.2s ease-in-out !important;
    }

    /* Efek saat kolom input diklik / fokus */
    div[data-baseweb="input"]:focus-within, 
    div[data-baseweb="select"]:focus-within > div {
        border-color: #0284c7 !important;
        box-shadow: 0 0 0 3px rgba(2, 132, 199, 0.2) !important;
    }

    /* Warna Teks Input */
    input, select, textarea {
        color: #0f172a !important;
        font-weight: 600 !important;
    }

    /* Label di atas Kolom Input */
    [data-labelledby] label, label {
        color: #1e293b !important;
        font-weight: 700 !important;
        font-size: 0.9rem !important;
    }

    /* Sidebar Terang & Clean */
    [data-testid="stSidebar"] {
        background-color: #f8fafc !important;
        border-right: 1px solid #e2e8f0 !important;
    }

    [data-testid="stSidebar"] * {
        color: #0f172a !important;
    }

    /* --- NAVIGASI SIDEBAR KARTU MODERN --- */
    [data-testid="stSidebar"] div[role="radiogroup"] {
        gap: 8px;
    }

    [data-testid="stSidebar"] div[role="radiogroup"] > label {
        background: #ffffff !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 12px !important;
        padding: 12px 16px !important;
        margin-bottom: 4px !important;
        box-shadow: 0 2px 4px rgba(0, 0, 0, 0.02) !important;
        transition: all 0.2s ease-in-out !important;
        cursor: pointer !important;
        width: 100% !important;
    }

    [data-testid="stSidebar"] div[role="radiogroup"] > label:hover {
        border-color: #0284c7 !important;
        background-color: #f0f9ff !important;
        transform: translateY(-1px);
    }

    [data-testid="stSidebar"] div[role="radiogroup"] > label:has(input:checked) {
        background: #e0f2fe !important;
        border: 2px solid #0284c7 !important;
        box-shadow: 0 4px 12px rgba(2, 132, 199, 0.15) !important;
    }

    [data-testid="stSidebar"] div[role="radiogroup"] label div[data-testid="stMarkdownContainer"] p {
        font-weight: 700 !important;
        color: #0f172a !important;
        font-size: 0.95rem !important;
    }

    /* Heading Text Styling */
    h1 {
        color: #0369a1 !important;
        font-weight: 800 !important;
        letter-spacing: -0.5px;
    }

    h2, h3 {
        color: #0284c7 !important;
        font-weight: 700 !important;
    }

    /* Kartu Metric Keuangan Serba Terang */
    [data-testid="stMetric"] {
        background: #ffffff;
        border: 1px solid #cbd5e1;
        border-top: 4px solid #0284c7;
        border-radius: 12px;
        padding: 16px 20px;
        box-shadow: 0 4px 10px rgba(0, 0, 0, 0.03);
    }

    [data-testid="stMetricValue"] {
        font-size: 1.4rem !important;
        font-weight: 800 !important;
        color: #0f172a !important;
    }

    [data-testid="stMetricLabel"] {
        font-size: 0.85rem !important;
        font-weight: 700 !important;
        color: #475569 !important;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }

    /* Tombol Utama Gradasi Emas-Cerah */
    .stButton > button {
        background: linear-gradient(135deg, #d97706 0%, #b45309 100%) !important;
        color: #ffffff !important;
        font-weight: 700 !important;
        font-size: 0.95rem !important;
        border-radius: 10px !important;
        padding: 12px 24px !important;
        border: none !important;
        box-shadow: 0 4px 14px rgba(217, 119, 6, 0.25) !important;
        transition: all 0.25s ease !important;
        width: 100% !important;
        margin-top: 10px !important;
    }

    .stButton > button:hover {
        background: linear-gradient(135deg, #f59e0b 0%, #d97706 100%) !important;
        box-shadow: 0 6px 18px rgba(217, 119, 6, 0.4) !important;
        transform: translateY(-1px);
    }

    /* Container Info Box */
    .stAlert {
        border-radius: 10px;
    }
    </style>
""", unsafe_allow_html=True)

# ==========================================
# 2. MANAGEMENT DATABASE SQLITE
# ==========================================
DB_FILE = "farm789.db"

def get_connection():
    conn = sqlite3.connect(DB_FILE, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_connection()
    cursor = conn.cursor()
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS transaksi (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tanggal TEXT NOT NULL,
            jenis_transaksi TEXT NOT NULL,
            jenis_ikan TEXT,
            ukuran TEXT,
            nama_item TEXT,
            qty REAL DEFAULT 1,
            satuan TEXT,
            harga_satuan REAL DEFAULT 0,
            total_nominal REAL NOT NULL,
            keterangan TEXT
        )
    """)
    
    cursor.execute("PRAGMA table_info(transaksi)")
    existing_columns = [col[1] for col in cursor.fetchall()]
    
    new_columns = {
        "jenis_ikan": "TEXT",
        "ukuran": "TEXT",
        "nama_item": "TEXT",
        "qty": "REAL DEFAULT 1",
        "satuan": "TEXT",
        "harga_satuan": "REAL DEFAULT 0"
    }
    
    for col_name, col_type in new_columns.items():
        if col_name not in existing_columns:
            cursor.execute(f"ALTER TABLE transaksi ADD COLUMN {col_name} {col_type}")
            
    conn.commit()
    conn.close()

init_db()

# ==========================================
# 3. HELPER FUNCTIONS (CRUD & FORMATTING)
# ==========================================
def load_data(start_date=None, end_date=None):
    conn = get_connection()
    query = "SELECT * FROM transaksi"
    params = []
    
    if start_date and end_date:
        query += " WHERE tanggal BETWEEN ? AND ?"
        params = [start_date.strftime("%Y-%m-%d"), end_date.strftime("%Y-%m-%d")]
        
    query += " ORDER BY tanggal ASC, id ASC"
    df = pd.read_sql_query(query, conn, params=params)
    conn.close()
    return df

def add_transaksi(tanggal, jenis_transaksi, jenis_ikan, ukuran, nama_item, qty, satuan, harga_satuan, total_nominal, keterangan):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO transaksi (tanggal, jenis_transaksi, jenis_ikan, ukuran, nama_item, qty, satuan, harga_satuan, total_nominal, keterangan)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (tanggal.strftime("%Y-%m-%d"), jenis_transaksi, jenis_ikan, ukuran, nama_item, qty, satuan, harga_satuan, total_nominal, keterangan))
    conn.commit()
    conn.close()

def update_transaksi(id_trans, tanggal, jenis_ikan, ukuran, nama_item, qty, satuan, harga_satuan, total_nominal, keterangan):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
        UPDATE transaksi 
        SET tanggal=?, jenis_ikan=?, ukuran=?, nama_item=?, qty=?, satuan=?, harga_satuan=?, total_nominal=?, keterangan=?
        WHERE id=?
    """, (tanggal.strftime("%Y-%m-%d"), jenis_ikan, ukuran, nama_item, qty, satuan, harga_satuan, total_nominal, keterangan, id_trans))
    conn.commit()
    conn.close()

def delete_transaksi(id_trans):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM transaksi WHERE id=?", (id_trans,))
    conn.commit()
    conn.close()

def fmt_rp(val):
    if val is None:
        return "Rp 0"
    if val < 0:
        return f"-Rp {abs(val):,.0f}".replace(",", ".")
    return f"Rp {val:,.0f}".replace(",", ".")

# Option Lists untuk Dropdown
IKAN_OPTIONS = ["-- Pilih Jenis Ikan --", "Discus", "Black Ghost", "Molly", "Guppy", "Angelfish", "Lainnya / Custom"]
UKURAN_OPTIONS = ["-- Pilih Ukuran --", "1.5 inch", "2 inch", "2.5 inch", "3 inch", "Indukan / Pair", "Custom / Lainnya"]
SATUAN_OPTIONS = ["pcs", "box", "set", "kaleng", "liter", "pack", "unit", "ekor/pair", "Custom / Lainnya"]

# Callback Functions untuk Simpan & Validasi
def submit_penjualan():
    ikan = st.session_state["sb_ikan"]
    ukuran = st.session_state["sb_ukuran"]
    qty = st.session_state.get("num_qty")
    harga = st.session_state.get("num_harga")
    
    if ikan == "-- Pilih Jenis Ikan --":
        st.session_state["msg_error"] = "⚠️ Harap pilih Jenis Ikan!"
        return
    if ukuran == "-- Pilih Ukuran --":
        st.session_state["msg_error"] = "⚠️ Harap pilih Ukuran / Size!"
        return
    if qty is None or harga is None:
        st.session_state["msg_error"] = "⚠️ Harap isi Jumlah (Qty) dan Harga Satuan!"
        return
        
    tgl = st.session_state["in_tgl"]
    ket = st.session_state["txt_ket"].strip()
    total = qty * harga
    
    add_transaksi(tgl, "IN_PENJUALAN", ikan, ukuran, "-", qty, "ekor/pair", harga, total, ket if ket else "-")
    st.session_state["msg_success"] = f"Berhasil menyimpan Penjualan Ikan sebesar {fmt_rp(total)}"
    st.session_state["sb_ikan"] = "-- Pilih Jenis Ikan --"
    st.session_state["sb_ukuran"] = "-- Pilih Ukuran --"
    st.session_state["num_qty"] = None
    st.session_state["num_harga"] = None
    st.session_state["txt_ket"] = ""

def submit_modal():
    total = st.session_state.get("num_modal")
    
    if total is None:
        st.session_state["msg_error"] = "⚠️ Harap isi Total Nominal Modal!"
        return
        
    tgl = st.session_state["in_modal_tgl"]
    ket = st.session_state["txt_modal"].strip()
    
    add_transaksi(tgl, "IN_MODAL", "-", "-", "Suntikan Modal", 1, "paket", total, total, ket if ket else "-")
    st.session_state["msg_success"] = f"Berhasil menyimpan Suntikan Modal sebesar {fmt_rp(total)}"
    st.session_state["num_modal"] = None
    st.session_state["txt_modal"] = ""

def submit_opex():
    nama = st.session_state["txt_item_out"].strip()
    qty = st.session_state.get("qty_out")
    harga = st.session_state.get("hrg_out")
    satuan_opt = st.session_state["sb_sat_out"]
    satuan_custom = st.session_state.get("txt_sat_out_custom", "").strip()
    
    satuan_final = satuan_custom if satuan_opt == "Custom / Lainnya" else satuan_opt
    
    if not nama:
        st.session_state["msg_error"] = "⚠️ Harap isi Nama Item / Barang!"
        return
    if qty is None or harga is None:
        st.session_state["msg_error"] = "⚠️ Harap isi Jumlah (Qty) dan Harga Satuan!"
        return
    if satuan_opt == "Custom / Lainnya" and not satuan_custom:
        st.session_state["msg_error"] = "⚠️ Harap isi Nama Satuan Custom!"
        return

    tgl = st.session_state["out_tgl"]
    ket = st.session_state["ket_out"].strip()
    total = qty * harga
    
    add_transaksi(tgl, "OUT_OPERASIONAL", "-", "-", nama, qty, satuan_final, harga, total, ket if ket else "-")
    st.session_state["msg_success"] = f"Berhasil menyimpan Pengeluaran sebesar {fmt_rp(total)}"
    st.session_state["txt_item_out"] = ""
    st.session_state["qty_out"] = None
    st.session_state["hrg_out"] = None
    st.session_state["txt_sat_out_custom"] = ""
    st.session_state["ket_out"] = ""

def submit_capex():
    nama = st.session_state["txt_aset"].strip()
    qty = st.session_state.get("qty_capex")
    harga = st.session_state.get("hrg_capex")
    satuan_opt = st.session_state["sb_sat_capex"]
    satuan_custom = st.session_state.get("txt_sat_capex_custom", "").strip()
    
    satuan_final = satuan_custom if satuan_opt == "Custom / Lainnya" else satuan_opt
    
    if not nama:
        st.session_state["msg_error"] = "⚠️ Harap isi Nama Aset!"
        return
    if qty is None or harga is None:
        st.session_state["msg_error"] = "⚠️ Harap isi Jumlah Qty dan Harga Per Unit!"
        return
    if satuan_opt == "Custom / Lainnya" and not satuan_custom:
        st.session_state["msg_error"] = "⚠️ Harap isi Nama Satuan Custom!"
        return

    tgl = st.session_state["capex_tgl"]
    ket = st.session_state["ket_capex"].strip()
    total = qty * harga
    
    add_transaksi(tgl, "CAPEX_ASET", "-", "-", nama, qty, satuan_final, harga, total, ket if ket else "-")
    st.session_state["msg_success"] = f"Berhasil menyimpan Aset sebesar {fmt_rp(total)}"
    st.session_state["txt_aset"] = ""
    st.session_state["qty_capex"] = None
    st.session_state["hrg_capex"] = None
    st.session_state["txt_sat_capex_custom"] = ""
    st.session_state["ket_capex"] = ""

# Function Export Excel
def generate_excel_report(df_rep, start_date, end_date):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    sub_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    total_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    big_total_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # 1. SHEET SUMMARY DASHBOARD
    ws_sum = wb.create_sheet(title="Summary Dashboard")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.cell(row=1, column=1, value="FARM 789 AQUARIUM - LAPORAN KEUANGAN EKSEKUTIF").font = title_font
    ws_sum.cell(row=2, column=1, value=f"Periode: {start_date.strftime('%d/%m/%Y')} s/d {end_date.strftime('%d/%m/%Y')}").font = Font(size=10, italic=True)
    
    omset = df_rep[df_rep['jenis_transaksi'] == 'IN_PENJUALAN']['total_nominal'].sum()
    opex = df_rep[df_rep['jenis_transaksi'] == 'OUT_OPERASIONAL']['total_nominal'].sum()
    laba_bersih = omset - opex
    capex = df_rep[df_rep['jenis_transaksi'] == 'CAPEX_ASET']['total_nominal'].sum()
    modal_in = df_rep[df_rep['jenis_transaksi'] == 'IN_MODAL']['total_nominal'].sum()
    saldo_akhir = (modal_in + omset) - (opex + capex)
    
    summary_data = [
        ["IN_PENJUALAN", "Total Omset Penjualan Ikan (+IN)", omset],
        ["OUT_OPERASIONAL", "Total Pengeluaran Operasional Rutin (-OUT)", opex],
        ["NET_PROFIT", "Laba / Rugi Murni Operasional", laba_bersih],
        ["CAPEX_ASET", "Total Belanja Aset & Indukan (-OUT CapEx)", capex],
        ["IN_MODAL", "Total Suntikan Modal Masuk (+IN Modal)", modal_in],
        ["SALDO_KAS", "Posisi Saldo Kas Bersih Tersedia", saldo_akhir]
    ]
    
    ws_sum.cell(row=4, column=1, value="Kode Indikator").font = header_font
    ws_sum.cell(row=4, column=1).fill = header_fill
    ws_sum.cell(row=4, column=2, value="Deskripsi Arus Kas").font = header_font
    ws_sum.cell(row=4, column=2).fill = header_fill
    ws_sum.cell(row=4, column=3, value="Total Nominal (Rp)").font = header_font
    ws_sum.cell(row=4, column=3).fill = header_fill
    
    for idx, row in enumerate(summary_data, start=5):
        c1 = ws_sum.cell(row=idx, column=1, value=row[0])
        c2 = ws_sum.cell(row=idx, column=2, value=row[1])
        c3 = ws_sum.cell(row=idx, column=3, value=row[2])
        c3.number_format = '"Rp "#,##0;[Red]"-Rp "#,##0'
        
        for c in [c1, c2, c3]:
            c.border = thin_border
            if row[0] == "SALDO_KAS":
                c.font = big_total_font
                c.fill = total_fill
            elif row[0] == "NET_PROFIT":
                c.font = bold_font
                c.fill = sub_fill

    # 2. SHEET BUKU KAS & MUTASI SALDO BERJALAN
    ws_kas = wb.create_sheet(title="Buku Kas & Mutasi Saldo")
    ws_kas.views.sheetView[0].showGridLines = True
    
    ws_kas.cell(row=1, column=1, value="BUKU KAS & MUTASI SALDO BERJALAN").font = title_font
    ws_kas.cell(row=2, column=1, value=f"Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}").font = Font(size=10, italic=True)
    
    headers_kas = ["Tanggal", "Jenis Transaksi", "Rincian Item / Keterangan", "Uang Masuk (+IN)", "Uang Keluar (-OUT)", "Saldo Berjalan (Rp)"]
    for c_idx, h_text in enumerate(headers_kas, start=1):
        cell = ws_kas.cell(row=4, column=c_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    row_idx = 5
    for r_idx, row in df_rep.iterrows():
        ws_kas.cell(row=row_idx, column=1, value=row['tanggal']).alignment = Alignment(horizontal="center")
        ws_kas.cell(row=row_idx, column=2, value=row['jenis_transaksi'])
        
        item_text = row['nama_item'] if row['nama_item'] != '-' else f"{row['jenis_ikan']} ({row['ukuran']})"
        if row['keterangan'] and row['keterangan'] != '-':
            item_text += f" - {row['keterangan']}"
        ws_kas.cell(row=row_idx, column=3, value=item_text)
        
        is_in = row['jenis_transaksi'] in ['IN_PENJUALAN', 'IN_MODAL']
        val_in = row['total_nominal'] if is_in else 0
        val_out = row['total_nominal'] if not is_in else 0
        
        c_in = ws_kas.cell(row=row_idx, column=4, value=val_in)
        c_out = ws_kas.cell(row=row_idx, column=5, value=val_out)
        c_in.number_format = '"Rp "#,##0'
        c_out.number_format = '"Rp "#,##0'
        
        c_saldo = ws_kas.cell(row=row_idx, column=6)
        if row_idx == 5:
            c_saldo.value = f"=D5-E5"
        else:
            c_saldo.value = f"=F{row_idx-1}+D{row_idx}-E{row_idx}"
        c_saldo.number_format = '"Rp "#,##0;[Red]"-Rp "#,##0'
        c_saldo.font = bold_font
        
        for col_num in range(1, 7):
            ws_kas.cell(row=row_idx, column=col_num).border = thin_border
        row_idx += 1

    # BARIS TOTAL SALDO KAS AKHIR
    ws_kas.cell(row=row_idx, column=1, value="TOTAL SISA SALDO KAS AKHIR").font = big_total_font
    ws_kas.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    
    c_tot_end = ws_kas.cell(row=row_idx, column=6)
    last_row_data = row_idx - 1
    c_tot_end.value = f"=F{last_row_data}" if last_row_data >= 5 else 0
    c_tot_end.number_format = '"Rp "#,##0;[Red]"-Rp "#,##0'
    c_tot_end.font = big_total_font
    
    for col_num in range(1, 7):
        c_elem = ws_kas.cell(row=row_idx, column=col_num)
        c_elem.fill = total_fill
        c_elem.border = thin_border
        if col_num == 1:
            c_elem.alignment = Alignment(horizontal="right")

    # Helper untuk Sheet Rinci Lainnya
    def add_detail_sheet(sheet_title, df_filtered, col_list):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        ws.cell(row=1, column=1, value=f"RINCIAN: {sheet_title.upper()}").font = title_font
        ws.cell(row=2, column=1, value=f"Periode: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}").font = Font(size=10, italic=True)
        
        for col_idx, col_name in enumerate(col_list, start=1):
            cell = ws.cell(row=4, column=col_idx, value=col_name.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
            
        r_start = 5
        for r_idx, row_data in df_filtered.iterrows():
            for c_idx, col_name in enumerate(col_list, start=1):
                val = row_data[col_name]
                cell = ws.cell(row=r_start, column=c_idx, value=val)
                cell.border = thin_border
                if col_name in ['harga_satuan', 'total_nominal']:
                    cell.number_format = '"Rp "#,##0'
                elif col_name == 'qty':
                    cell.number_format = '#,##0.00'
                elif col_name == 'tanggal':
                    cell.alignment = Alignment(horizontal="center")
            r_start += 1
            
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    df_lr = df_rep[df_rep['jenis_transaksi'].isin(['IN_PENJUALAN', 'OUT_OPERASIONAL'])]
    add_detail_sheet("Laba Rugi Operasional", df_lr, ['tanggal', 'jenis_transaksi', 'jenis_ikan', 'ukuran', 'nama_item', 'qty', 'satuan', 'harga_satuan', 'total_nominal', 'keterangan'])
    
    df_capex = df_rep[df_rep['jenis_transaksi'] == 'CAPEX_ASET']
    add_detail_sheet("Belanja Aset CapEx", df_capex, ['tanggal', 'nama_item', 'qty', 'satuan', 'harga_satuan', 'total_nominal', 'keterangan'])
    
    df_modal = df_rep[df_rep['jenis_transaksi'] == 'IN_MODAL']
    add_detail_sheet("Suntikan Modal", df_modal, ['tanggal', 'nama_item', 'total_nominal', 'keterangan'])

    # Adjust Column Widths
    for ws_curr in [ws_sum, ws_kas]:
        for col in ws_curr.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_curr.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

# ==========================================
# 4. SIDEBAR NAVIGASI
# ==========================================
if os.path.exists("logo.png"):
    st.sidebar.image("logo.png", width="stretch")
elif os.path.exists("logo.jpg"):
    st.sidebar.image("logo.jpg", width="stretch")

st.sidebar.caption("Sistem Keuangan Budidaya Ikan Hias")

menu = st.sidebar.radio(
    "Pilih Menu Laporan:",
    [
        "🏠 Dashboard Utama",
        "⚡ Form Input Transaksi",
        "🧮 Kalkulator HPP",
        "🖨️ Laporan & Export"
    ]
)

st.sidebar.markdown("---")
st.sidebar.subheader("📅 Filter Periode Laporan")
today = datetime.date.today()
first_day = today.replace(day=1)

start_date = st.sidebar.date_input("Dari Tanggal", first_day, format="DD/MM/YYYY", key="filter_start")
end_date = st.sidebar.date_input("Sampai Tanggal", today, format="DD/MM/YYYY", key="filter_end")

# ==========================================
# DIALOG / MODAL POPUP EDIT TRANSAKSI
# ==========================================
@st.dialog("✏️ Edit Transaksi Keuangan")
def edit_dialog(row):
    tgl_dt = datetime.datetime.strptime(row['tanggal'], "%Y-%m-%d").date()
    edit_tgl = st.date_input("📅 Tanggal", tgl_dt, format="DD/MM/YYYY")
    
    if row['jenis_transaksi'] == "IN_PENJUALAN":
        list_ikan = IKAN_OPTIONS[1:]
        idx_ikan = list_ikan.index(row['jenis_ikan']) if row['jenis_ikan'] in list_ikan else 0
        edit_ikan = st.selectbox("Jenis Ikan", list_ikan, index=idx_ikan)
        
        list_ukuran = UKURAN_OPTIONS[1:]
        idx_ukuran = list_ukuran.index(row['ukuran']) if row['ukuran'] in list_ukuran else 0
        edit_ukuran = st.selectbox("Ukuran / Size", list_ukuran, index=idx_ukuran)
        
        edit_qty = st.number_input("Jumlah (Qty)", min_value=1, value=int(row['qty']), step=1, format="%d")
        edit_harga = st.number_input("Harga Satuan (Rp)", min_value=0, value=int(row['harga_satuan']), step=1000, format="%d")
        total_calc = edit_qty * edit_harga
        st.markdown(f"**Total Nominal:** {fmt_rp(total_calc)}")
        edit_ket = st.text_input("Keterangan", value="" if row['keterangan'] == "-" else row['keterangan'], placeholder="Contoh: Red Melon 2 inch - Pak Budi")
        
        if st.button("💾 SIMPAN PERUBAHAN", key="btn_save_edit"):
            update_transaksi(row['id'], edit_tgl, edit_ikan, edit_ukuran, "-", edit_qty, "ekor/pair", edit_harga, total_calc, edit_ket if edit_ket else "-")
            st.success("Berhasil memperbarui transaksi!")
            st.rerun()

    elif row['jenis_transaksi'] == "IN_MODAL":
        edit_total = st.number_input("Total Modal (Rp)", min_value=0, value=int(row['total_nominal']), step=10000, format="%d")
        edit_ket = st.text_input("Keterangan Sumber Dana", value="" if row['keterangan'] == "-" else row['keterangan'], placeholder="Contoh: Tabungan Pribadi, Pinjaman Bank")
        
        if st.button("💾 SIMPAN PERUBAHAN", key="btn_save_edit_m"):
            update_transaksi(row['id'], edit_tgl, "-", "-", "Suntikan Modal", 1, "paket", edit_total, edit_total, edit_ket if edit_ket else "-")
            st.success("Berhasil memperbarui modal!")
            st.rerun()

    else:
        edit_nama = st.text_input("Nama Item / Barang", value="" if row['nama_item'] == "-" else row['nama_item'], placeholder="Contoh: Tank Aquarium 100x40x40")
        edit_qty = st.number_input("Jumlah (Qty)", min_value=0.1, value=float(row['qty']), step=0.1, format="%g")
        
        sat_val = row['satuan'] if row['satuan'] in SATUAN_OPTIONS else "Custom / Lainnya"
        idx_sat = SATUAN_OPTIONS.index(sat_val) if sat_val in SATUAN_OPTIONS else 0
        edit_sat_opt = st.selectbox("Satuan", SATUAN_OPTIONS, index=idx_sat)
        edit_sat_custom = st.text_input("Ketik Satuan Custom", value=row['satuan'] if sat_val == "Custom / Lainnya" else "") if edit_sat_opt == "Custom / Lainnya" else ""
        
        edit_sat_final = edit_sat_custom if edit_sat_opt == "Custom / Lainnya" else edit_sat_opt
        
        edit_harga = st.number_input("Harga Satuan (Rp)", min_value=0, value=int(row['harga_satuan']), step=1000, format="%d")
        total_calc = edit_qty * edit_harga
        st.markdown(f"**Total Nominal:** {fmt_rp(total_calc)}")
        edit_ket = st.text_input("Catatan Tambahan", value="" if row['keterangan'] == "-" else row['keterangan'], placeholder="Contoh: Catatan opsional")
        
        if st.button("💾 SIMPAN PERUBAHAN", key="btn_save_edit_o"):
            update_transaksi(row['id'], edit_tgl, "-", "-", edit_nama if edit_nama else "-", edit_qty, edit_sat_final if edit_sat_final else "unit", edit_harga, total_calc, edit_ket if edit_ket else "-")
            st.success("Berhasil memperbarui transaksi!")
            st.rerun()

# ==========================================
# FRAME 1: DASHBOARD UTAMA
# ==========================================
if menu == "🏠 Dashboard Utama":
    st.title("📊 Dashboard Keuangan Farm")
    st.caption(f"Periode Terfilter: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}")
    
    df = load_data(start_date, end_date)
    df_all = load_data()
    
    omset_penjualan = df[df['jenis_transaksi'] == 'IN_PENJUALAN']['total_nominal'].sum()
    total_operasional = df[df['jenis_transaksi'] == 'OUT_OPERASIONAL']['total_nominal'].sum()
    laba_bersih = omset_penjualan - total_operasional
    margin = (laba_bersih / omset_penjualan * 100) if omset_penjualan > 0 else 0
    
    total_aset_all = df_all[df_all['jenis_transaksi'] == 'CAPEX_ASET']['total_nominal'].sum()
    total_modal_all = df_all[df_all['jenis_transaksi'] == 'IN_MODAL']['total_nominal'].sum()
    total_omset_all = df_all[df_all['jenis_transaksi'] == 'IN_PENJUALAN']['total_nominal'].sum()
    total_operasional_all = df_all[df_all['jenis_transaksi'] == 'OUT_OPERASIONAL']['total_nominal'].sum()
    
    total_laba_kumulatif = total_omset_all - total_operasional_all
    target_bep = max(total_aset_all, total_modal_all)
    persen_bep = (total_laba_kumulatif / target_bep * 100) if target_bep > 0 else 0
    persen_bep = min(persen_bep, 100.0)
    sisa_bep = max(0.0, target_bep - total_laba_kumulatif)

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("TOTAL OMSET (IN)", fmt_rp(omset_penjualan), "Penjualan Ikan")
    col2.metric("OPERASIONAL (OUT)", fmt_rp(total_operasional), "Biaya Rutin Pakan/Air")
    col3.metric("LABA BERSIH (NET)", fmt_rp(laba_bersih), f"Margin: {margin:.1f}%")
    col4.metric("BALIK MODAL (BEP)", f"{persen_bep:.1f}%", f"Sisa: {fmt_rp(sisa_bep)}")

    st.markdown("---")

    st.subheader("💡 Alokasi Pemisahan Kas (Otomatis Split Laba)")
    pos_hpp = total_operasional
    laba_positif = max(0.0, laba_bersih)
    pos_bep = laba_positif * 0.50
    pos_pengembangan = laba_positif * 0.30
    pos_owner = laba_positif * 0.20

    p1, p2, p3, p4 = st.columns(4)
    p1.info(f"**Putaran HPP Operasional**\n### {fmt_rp(pos_hpp)}\n*(Untuk Pakan & Listrik)*")
    p2.success(f"**Tabungan BEP (50%)**\n### {fmt_rp(pos_bep)}\n*(Pengembalian Investasi)*")
    p3.warning(f"**Kas Pengembangan (30%)**\n### {fmt_rp(pos_pengembangan)}\n*(Ekspansi Tank & Indukan)*")
    p4.error(f"**Profit Owner / Gaji (20%)**\n### {fmt_rp(pos_owner)}\n*(Aman Ditarik Pribadi)*")

    st.markdown("---")

    g1, g2 = st.columns([6, 4])
    with g1:
        st.subheader("📊 Tren Penjualan per Jenis Ikan")
        df_sales = df[df['jenis_transaksi'] == 'IN_PENJUALAN']
        if not df_sales.empty:
            fig_sales = px.bar(
                df_sales, 
                x='jenis_ikan', 
                y='total_nominal', 
                color='ukuran',
                title="Penjualan Berdasarkan Ikan & Ukuran",
                labels={'jenis_ikan': 'Jenis Ikan', 'total_nominal': 'Total (Rp)'},
                color_discrete_sequence=px.colors.qualitative.Bold
            )
            fig_sales.update_layout(paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
            st.plotly_chart(fig_sales, width="stretch")
        else:
            st.info("Belum ada data penjualan pada periode ini.")

    with g2:
        st.subheader("🥧 Pengeluaran Terboros (Top OUT Item)")
        df_out = df[df['jenis_transaksi'] == 'OUT_OPERASIONAL']
        if not df_out.empty:
            top_out = df_out.groupby('nama_item')['total_nominal'].sum().reset_index()
            top_out = top_out.sort_values(by='total_nominal', ascending=False)
            
            fig_out = px.pie(
                top_out, 
                names='nama_item', 
                values='total_nominal', 
                title="Porsi Biaya Operasional",
                color_discrete_sequence=px.colors.sequential.Agsunset
            )
            fig_out.update_layout(paper_bgcolor='rgba(0,0,0,0)')
            st.plotly_chart(fig_out, width="stretch")
        else:
            st.info("Belum ada data pengeluaran operasional.")

    st.markdown("---")
    st.subheader("🏛️ Ringkasan Arus Kas & Modal Keseluruhan (All-Time)")
    
    c_m1, c_m2, c_m3, c_m4 = st.columns(4)
    c_m1.metric("💰 Total Modal Masuk", fmt_rp(total_modal_all), "Investasi Modal")
    c_m2.metric("🏗️ Belanja Aset (CapEx)", fmt_rp(total_aset_all), "Tank, Rak, Indukan")
    c_m3.metric("🐟 Total Penjualan", fmt_rp(total_omset_all), "Omset Penjualan")
    c_m4.metric("💸 Total Operasional", fmt_rp(total_operasional_all), "Biaya Rutin")
    
    total_uang_masuk_all = total_modal_all + total_omset_all
    total_uang_keluar_all = total_operasional_all + total_aset_all
    kas_tersedia = total_uang_masuk_all - total_uang_keluar_all
    
    st.markdown(f"""
        <div style="background-color: #f0f9ff; border: 2px solid #0284c7; border-left: 8px solid #d97706; padding: 20px; border-radius: 12px; margin-top: 15px;">
            <p style="margin: 0; font-size: 0.9rem; color: #0369a1; font-weight: 700; text-transform: uppercase;">💵 Posisi Kas Murni Tersedia</p>
            <h2 style="margin: 5px 0 0 0; color: #b45309; font-size: 2.2rem; font-weight: 800;">{fmt_rp(kas_tersedia)}</h2>
        </div>
    """, unsafe_allow_html=True)

# ==========================================
# FRAME 2: FORM INPUT TRANSAKSI
# ==========================================
elif menu == "⚡ Form Input Transaksi":
    st.title("⚡ Form Input Transaksi Cepat")
    
    if "msg_success" in st.session_state:
        st.success(st.session_state["msg_success"])
        del st.session_state["msg_success"]
        
    if "msg_error" in st.session_state:
        st.error(st.session_state["msg_error"])
        del st.session_state["msg_error"]
    
    tab_pilihan = st.radio(
        "Pilih Jenis Transaksi:",
        ["💰 Uang Masuk (+IN)", "💸 Operasional (-OUT)", "📦 Belanja Aset (CapEx)"],
        horizontal=True,
        key="active_tab_form"
    )
    
    st.markdown("---")
    
    if tab_pilihan == "💰 Uang Masuk (+IN)":
        st.subheader("Pencatatan Uang Masuk")
        
        c_tipe1, c_tipe2 = st.columns([35, 65])
        with c_tipe1:
            tipe_in = st.selectbox("Tipe Uang Masuk", ["Penjualan Ikan", "Tambah Modal / Suntikan Dana"], key="select_tipe_in")
        
        st.markdown("<br>", unsafe_allow_html=True)
        
        if tipe_in == "Penjualan Ikan":
            col_l, col_r = st.columns(2)
            with col_l:
                st.date_input("📅 Tanggal Transaksi", datetime.date.today(), format="DD/MM/YYYY", key="in_tgl")
                st.selectbox("Jenis Ikan *", IKAN_OPTIONS, key="sb_ikan")
                st.selectbox("Ukuran / Size *", UKURAN_OPTIONS, key="sb_ukuran")
            
            with col_r:
                qty = st.number_input("Jumlah (Qty Ekor/Pair) *", min_value=1, value=None, step=1, format="%d", placeholder="Contoh: 10", key="num_qty")
                harga_satuan = st.number_input("Harga Satuan (Rp) *", min_value=0, value=None, step=1000, format="%d", placeholder="Contoh: 15000", key="num_harga")
                tot_val = (qty * harga_satuan) if (qty is not None and harga_satuan is not None) else 0
                st.markdown(f"### Total Auto: **{fmt_rp(tot_val)}**")
                st.text_input("Keterangan / Varietas & Nama Pembeli", value="", placeholder="Contoh: Red Melon 2 inch - Pak Budi", key="txt_ket")
                
            st.button("✓ SIMPAN PENJUALAN IKAN", key="btn_in", on_click=submit_penjualan)

        else:
            col_l, col_r = st.columns(2)
            with col_l:
                st.date_input("📅 Tanggal Transaksi", datetime.date.today(), format="DD/MM/YYYY", key="in_modal_tgl")
                st.number_input("Total Nominal Modal (Rp) *", min_value=0, value=None, step=50000, format="%d", placeholder="Contoh: 1000000", key="num_modal")
            
            with col_r:
                st.text_input("Keterangan Sumber Dana", value="", placeholder="Contoh: Tabungan Pribadi, Pinjaman Bank", key="txt_modal")
                
            st.button("✓ SIMPAN SUNTIKAN MODAL", key="btn_modal", on_click=submit_modal)

    elif tab_pilihan == "💸 Operasional (-OUT)":
        st.subheader("Pencatatan Pengeluaran Operasional (OUT)")
        col_o1, col_o2 = st.columns(2)
        
        with col_o1:
            st.date_input("📅 Tanggal Transaksi", datetime.date.today(), format="DD/MM/YYYY", key="out_tgl")
            st.text_input("Nama Item / Barang *", value="", placeholder="Contoh: Artemia, Cacing Sutra, Listrik", key="txt_item_out")
            qty_out = st.number_input("Jumlah (Qty) *", min_value=0.1, value=None, step=1.0, format="%g", placeholder="Contoh: 5", key="qty_out")
            
        with col_o2:
            sb_sat_out = st.selectbox("Satuan *", SATUAN_OPTIONS, key="sb_sat_out")
            if sb_sat_out == "Custom / Lainnya":
                st.text_input("Ketik Satuan Custom *", value="", placeholder="Contoh: botol / roll", key="txt_sat_out_custom")
            
            harga_out = st.number_input("Harga Satuan (Rp) *", min_value=0, value=None, step=5000, format="%d", placeholder="Contoh: 100000", key="hrg_out")
            tot_out_val = (qty_out * harga_out) if (qty_out is not None and harga_out is not None) else 0
            st.markdown(f"### Total Biaya Auto: **{fmt_rp(tot_out_val)}**")
            st.text_input("Catatan Tambahan", value="", placeholder="Contoh: Pembelian pakan burayak", key="ket_out")
        
        st.button("✓ SIMPAN PENGELUARAN", key="btn_out", on_click=submit_opex)

    elif tab_pilihan == "📦 Belanja Aset (CapEx)":
        st.subheader("Pencatatan Belanja Aset (CapEx)")
        col_c1, col_c2 = st.columns(2)
        
        with col_c1:
            st.date_input("📅 Tanggal Transaksi", datetime.date.today(), format="DD/MM/YYYY", key="capex_tgl")
            st.text_input("Nama Aset *", value="", placeholder="Contoh: Tank 100x40x40 cm, Indukan Pair", key="txt_aset")
            qty_capex = st.number_input("Qty Unit *", min_value=1.0, value=None, step=1.0, format="%g", placeholder="Contoh: 1", key="qty_capex")
            
        with col_c2:
            sb_sat_capex = st.selectbox("Satuan Aset *", SATUAN_OPTIONS, index=6, key="sb_sat_capex")
            if sb_sat_capex == "Custom / Lainnya":
                st.text_input("Ketik Satuan Custom *", value="", placeholder="Contoh: unit / set", key="txt_sat_capex_custom")
                
            harga_capex = st.number_input("Harga Per Unit (Rp) *", min_value=0, value=None, step=10000, format="%d", placeholder="Contoh: 500000", key="hrg_capex")
            tot_capex_val = (qty_capex * harga_capex) if (qty_capex is not None and harga_capex is not None) else 0
            st.markdown(f"### Total Investasi Aset: **{fmt_rp(tot_capex_val)}**")
            st.text_input("Spesifikasi / Catatan Aset", value="", placeholder="Contoh: Tank Baru / Indukan Red Melon", key="ket_capex")
        
        st.button("✓ SIMPAN BELANJA ASET", key="btn_capex", on_click=submit_capex)

    st.markdown("---")
    st.subheader("📋 Riwayat Transaksi & Fitur Edit / Delete")
    
    df_history = load_data(start_date, end_date)
    if not df_history.empty:
        for idx, row in df_history.iterrows():
            item_display = row['nama_item'] if row['nama_item'] != '-' else f"{row['jenis_ikan']} ({row['ukuran']})"
            with st.expander(f"📌 [{row['tanggal']}] {row['jenis_transaksi']} - {item_display} ({fmt_rp(row['total_nominal'])})"):
                c_edit1, c_edit2, c_edit3 = st.columns([7, 1.5, 1.5])
                with c_edit1:
                    st.write(f"**Detail:** Qty {row['qty']:g} {row['satuan']} @ {fmt_rp(row['harga_satuan'])}")
                    st.write(f"**Keterangan:** {row['keterangan'] if row['keterangan'] else '-'}")
                with c_edit2:
                    if st.button("✏️ Edit", key=f"edit_btn_{row['id']}"):
                        edit_dialog(row)
                with c_edit3:
                    if st.button("🗑️ Hapus", key=f"del_{row['id']}"):
                        delete_transaksi(row['id'])
                        st.success("Transaksi dihapus!")
                        st.rerun()
    else:
        st.info("Belum ada data transaksi di periode ini.")

# ==========================================
# FRAME 3: KALKULATOR HPP
# ==========================================
elif menu == "🧮 Kalkulator HPP":
    st.title("🧮 Kalkulator HPP & Formulasi Harga Jual")
    st.caption("Hitung modal murni per ekor/per batch pembesaran agar tidak salah pasang harga jual.")
    
    col_hpp1, col_hpp2 = st.columns(2)
    
    with col_hpp1:
        st.subheader("Input Batch Pembesaran")
        qty_batch = st.number_input("Jumlah Awal Burayak/Ikan (Ekor)", min_value=1, value=100, step=1, format="%d", key="batch_qty")
        sr_rate = st.number_input("Tingkat Kelangsungan Hidup / Survival Rate (%)", min_value=1.0, max_value=100.0, value=70.0, step=5.0, format="%.0f", key="batch_sr")
        
        qty_hidup = max(1, int(qty_batch * (sr_rate / 100.0)))
        st.caption(f"ℹ️ Estimasi ikan hidup sampai panen: **{qty_hidup} ekor** dari {qty_batch} ekor (Penyusutan {100-int(sr_rate)}%)")
        
        lama_rawat = st.number_input("Lama Perawatan (Hari)", min_value=1, value=60, step=1, format="%d", key="batch_hari")
        
        df_opex = load_data()
        df_opex_out = df_opex[df_opex['jenis_transaksi'] == 'OUT_OPERASIONAL']
        total_out_all = df_opex_out['total_nominal'].sum()
        
        biaya_harian_default = int((total_out_all / 30 / 10)) if total_out_all > 0 else 5000
        
        c_b1, c_b2 = st.columns([6, 4])
        with c_b1:
            biaya_input = st.number_input("Estimasi Biaya Routine per Tank (Rp)", min_value=0, value=biaya_harian_default, step=500, format="%d", key="batch_biaya_input")
        with c_b2:
            periode_biaya = st.selectbox("Periode Biaya", ["Per Hari", "Per Minggu", "Per Bulan"], key="batch_periode")
        
        if periode_biaya == "Per Minggu":
            biaya_per_tank_hari = biaya_input / 7.0
        elif periode_biaya == "Per Bulan":
            biaya_per_tank_hari = biaya_input / 30.0
        else:
            biaya_per_tank_hari = float(biaya_input)
            
        st.caption(f"💡 Konversi Biaya Rutin Harian: **{fmt_rp(biaya_per_tank_hari)} / hari**")
        modal_indukan = st.number_input("Beban Susut Indukan per Batch (Rp)", min_value=0, value=50000, step=5000, format="%d", key="batch_indukan")

    with col_hpp2:
        st.subheader("Hasil Kalkulasi Otomatis")
        total_biaya_pakan_listrik = lama_rawat * biaya_per_tank_hari
        total_hpp_batch = total_biaya_pakan_listrik + modal_indukan
        hpp_per_ekor = total_hpp_batch / qty_hidup
        
        st.info(f"**Total Biaya Rawat ({lama_rawat} Hari):** {fmt_rp(total_biaya_pakan_listrik)}")
        st.info(f"**Total HPP Modal 1 Batch:** {fmt_rp(total_hpp_batch)}")
        st.success(f"### HPP Modal Murni: **{fmt_rp(hpp_per_ekor)} / ekor**\n*(Dihitung dari {qty_hidup} ekor ikan yang hidup)*")
        
        st.markdown("---")
        st.subheader("💡 Rekomendasi Harga Jual Minimum")
        partai = hpp_per_ekor * 2.0
        eceran = hpp_per_ekor * 5.0
        
        st.write(f"• **Harga Partai / Borongan (Margin 100%):** **{fmt_rp(partai)} / ekor**")
        st.write(f"• **Harga Eceran / Selected Pair (Margin 400%):** **{fmt_rp(eceran * 2)} / pair**")

# ==========================================
# FRAME 4: LAPORAN & EXPORT
# ==========================================
elif menu == "🖨️ Laporan & Export":
    st.title("🖨️ Laporan Keuangan Otomatis & Terstruktur")
    st.caption(f"Periode Laporan: {start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}")
    
    df_rep = load_data(start_date, end_date)
    
    if not df_rep.empty:
        omset = df_rep[df_rep['jenis_transaksi'] == 'IN_PENJUALAN']['total_nominal'].sum()
        opex = df_rep[df_rep['jenis_transaksi'] == 'OUT_OPERASIONAL']['total_nominal'].sum()
        laba_bersih = omset - opex
        capex = df_rep[df_rep['jenis_transaksi'] == 'CAPEX_ASET']['total_nominal'].sum()
        modal_in = df_rep[df_rep['jenis_transaksi'] == 'IN_MODAL']['total_nominal'].sum()
        
        st.subheader("📌 Ringkasan Eksekutif Periode Ini")
        m_row1_1, m_row1_2, m_row1_3 = st.columns(3)
        m_row1_1.metric("💰 Omset Penjualan", fmt_rp(omset))
        m_row1_2.metric("💸 Biaya Operasional", fmt_rp(opex))
        m_row1_3.metric("📈 Laba / Rugi Murni", fmt_rp(laba_bersih))
        
        st.markdown("<br>", unsafe_allow_html=True)
        m_row2_1, m_row2_2, m_row2_3 = st.columns(3)
        m_row2_1.metric("🏗️ Belanja Aset (CapEx)", fmt_rp(capex))
        m_row2_2.metric("🏦 Suntikan Modal Masuk", fmt_rp(modal_in))
        
        st.markdown("---")
        
        tab_rep1, tab_rep2, tab_rep3, tab_rep4 = st.tabs([
            "📊 1. Laporan Laba Rugi", 
            "🏗️ 2. Aset & Inventaris", 
            "💳 3. Arus Kas & Modal",
            "📋 4. Buku Kas (Mutasi Saldo)"
        ])
        
        def style_df(data_frame):
            df_styled = data_frame.copy()
            if 'harga_satuan' in df_styled.columns:
                df_styled['harga_satuan'] = df_styled['harga_satuan'].apply(fmt_rp)
            if 'total_nominal' in df_styled.columns:
                df_styled['total_nominal'] = df_styled['total_nominal'].apply(fmt_rp)
            return df_styled

        with tab_rep1:
            st.subheader("📊 Laporan Laba / Rugi Murni Operasional")
            
            st.markdown("### 🟢 1. Pendapatan (Uang Masuk Penjualan Ikan)")
            df_in = df_rep[df_rep['jenis_transaksi'] == 'IN_PENJUALAN'][['tanggal', 'jenis_ikan', 'ukuran', 'qty', 'harga_satuan', 'total_nominal', 'keterangan']]
            if not df_in.empty:
                st.dataframe(style_df(df_in), width="stretch")
            else:
                st.info("Belum ada pendapatan penjualan di periode ini.")
            
            st.markdown("---")
            
            st.markdown("### 🔴 2. Beban Operasional (Pengeluaran Rutin Pakan/Listrik)")
            df_out = df_rep[df_rep['jenis_transaksi'] == 'OUT_OPERASIONAL'][['tanggal', 'nama_item', 'qty', 'satuan', 'harga_satuan', 'total_nominal', 'keterangan']]
            if not df_out.empty:
                st.dataframe(style_df(df_out), width="stretch")
            else:
                st.info("Belum ada pengeluaran operasional di periode ini.")

        with tab_rep2:
            st.subheader("🏗️ Laporan Belanja Aset & Indukan (CapEx)")
            df_capex = df_rep[df_rep['jenis_transaksi'] == 'CAPEX_ASET'][['tanggal', 'nama_item', 'qty', 'satuan', 'harga_satuan', 'total_nominal', 'keterangan']]
            if not df_capex.empty:
                st.dataframe(style_df(df_capex), width="stretch")
            else:
                st.info("Tidak ada belanja aset pada periode ini.")

        with tab_rep3:
            st.subheader("💳 Laporan Arus Kas & Suntikan Modal")
            df_modal = df_rep[df_rep['jenis_transaksi'] == 'IN_MODAL'][['tanggal', 'nama_item', 'total_nominal', 'keterangan']]
            if not df_modal.empty:
                st.dataframe(style_df(df_modal), width="stretch")
            else:
                st.info("Tidak ada suntikan modal pada periode ini.")

        with tab_rep4:
            st.subheader("📋 Buku Kas & Mutasi Saldo Berjalan")
            
            df_kas_ui = df_rep.copy()
            saldo_run = 0
            s_list = []
            in_list = []
            out_list = []
            
            for idx, r in df_kas_ui.iterrows():
                is_in = r['jenis_transaksi'] in ['IN_PENJUALAN', 'IN_MODAL']
                val_in = r['total_nominal'] if is_in else 0
                val_out = r['total_nominal'] if not is_in else 0
                
                saldo_run += (val_in - val_out)
                in_list.append(fmt_rp(val_in) if val_in > 0 else "-")
                out_list.append(fmt_rp(val_out) if val_out > 0 else "-")
                s_list.append(fmt_rp(saldo_run))
                
            df_kas_ui['Uang Masuk (+IN)'] = in_list
            df_kas_ui['Uang Keluar (-OUT)'] = out_list
            df_kas_ui['Saldo Berjalan'] = s_list
            
            bg_color = "#f0fdf4" if saldo_run >= 0 else "#fef2f2"
            border_color = "#16a34a" if saldo_run >= 0 else "#dc2626"
            text_color = "#15803d" if saldo_run >= 0 else "#b91c1c"
            status_label = "SURPLUS (SISA KAS TERSEDIA)" if saldo_run >= 0 else "DEFISIT (SALDO MINUS)"
            
            st.markdown(f"""
                <div style="background-color: {bg_color}; border: 2px solid {border_color}; padding: 18px 24px; border-radius: 10px; margin-bottom: 20px;">
                    <p style="margin: 0; font-size: 0.95rem; font-weight: 700; color: #475569; letter-spacing: 0.5px;">
                        💵 TOTAL POSISI SALDO KAS PERIODE INI ({status_label}):
                    </p>
                    <h2 style="margin: 5px 0 0 0; font-size: 2.2rem; font-weight: 800; color: {text_color};">
                        {fmt_rp(saldo_run)}
                    </h2>
                </div>
            """, unsafe_allow_html=True)
            
            disp_cols = ['tanggal', 'jenis_transaksi', 'nama_item', 'Uang Masuk (+IN)', 'Uang Keluar (-OUT)', 'Saldo Berjalan', 'keterangan']
            st.dataframe(df_kas_ui[disp_cols], width="stretch")

        st.markdown("---")
        
        if OPENPYXL_AVAILABLE:
            excel_bytes = generate_excel_report(df_rep, start_date, end_date)
            st.download_button(
                label="📊 DOWNLOAD LAPORAN EXCEL LENGKAP WITH FORMULA SALDO (.XLSX)",
                data=excel_bytes,
                file_name=f"Laporan_Keuangan_FARM789_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            st.warning("Gunakan `pip install openpyxl` untuk mengaktifkan fitur cetak laporan Excel profesional.")